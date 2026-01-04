from datetime import date

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_POST
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db.models import Q, Sum
from django.http import HttpResponse

from .forms import PendaftaranForm
from .forms_admin import (
    AdminPendaftaranForm,
    PembayaranDaftarUlangForm,
)
from .models import (
    Pendaftaran,
    PembayaranDaftarUlang,
    LogAktivitas,
)
from .permissions import role_required

# =========================
# QR CODE
# =========================
import qrcode
from io import BytesIO
import base64

# =========================
# EXCEL
# =========================
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


# =====================================================
# UTIL
# =====================================================
def generate_qr_base64(data):
    qr = qrcode.QRCode(
        version=3,
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=12,
        border=4,
    )
    qr.add_data(data)
    qr.make(fit=True)

    img = qr.make_image(fill_color="black", back_color="white")

    buffer = BytesIO()
    img.save(buffer, format="PNG")
    img_str = base64.b64encode(buffer.getvalue()).decode()

    return f"data:image/png;base64,{img_str}"


# =====================================================
# PUBLIC
# =====================================================
def home(request):
    if request.method == 'POST':
        form = PendaftaranForm(request.POST)
        if form.is_valid():
            pendaftaran = form.save()
            messages.success(request, 'Pendaftaran berhasil!')
            return redirect('sukses', pk=pendaftaran.pk)
    else:
        form = PendaftaranForm()

    return render(request, 'pendaftaran.html', {'form': form})


def sukses(request, pk):
    pendaftaran = get_object_or_404(Pendaftaran, pk=pk)
    return render(request, 'sukses.html', {'pendaftaran': pendaftaran})


def print_kartu(request, nomor_pendaftaran):
    pendaftaran = get_object_or_404(
        Pendaftaran,
        nomor_pendaftaran=nomor_pendaftaran
    )

    qr_data = request.build_absolute_uri(
        f"/kartu/{pendaftaran.nomor_pendaftaran}/"
    )
    qr_image = generate_qr_base64(qr_data)

    return render(request, 'kartu_pendaftaran.html', {
        'pendaftaran': pendaftaran,
        'qr_image': qr_image,
    })


# =====================================================
# ADMIN PANEL
# =====================================================
@login_required
@role_required('SUPERADMIN', 'ADMIN', 'PANITIA')
def dashboard_admin(request):
    qs = Pendaftaran.objects.all()

    return render(request, 'adminpanel/dashboard.html', {
        'pendaftarans': qs.order_by('-tanggal_pendaftaran')[:10],
        'total': qs.count(),
        'terdaftar': qs.filter(status='terdaftar').count(),
        'diterima': qs.filter(status='diterima').count(),
        'ditolak': qs.filter(status='ditolak').count(),
    })


@login_required
@role_required('SUPERADMIN', 'ADMIN', 'PANITIA')
def admin_pendaftaran_list(request):
    status_filter = request.GET.get('status')
    jurusan_filter = request.GET.get('jurusan')
    keyword = request.GET.get('q')

    qs = Pendaftaran.objects.all().order_by('-tanggal_pendaftaran')

    if status_filter:
        qs = qs.filter(status=status_filter)

    if jurusan_filter:
        qs = qs.filter(jurusan=jurusan_filter)

    if keyword:
        qs = qs.filter(
            Q(nama_lengkap__icontains=keyword) |
            Q(nik__icontains=keyword) |
            Q(nomor_pendaftaran__icontains=keyword)
        )

    paginator = Paginator(qs, 10)
    page_obj = paginator.get_page(request.GET.get('page'))

    return render(request, 'adminpanel/pendaftaran_list.html', {
        'pendaftarans': page_obj,
        'status_filter': status_filter,
        'jurusan_filter': jurusan_filter,
        'keyword': keyword,
        'page_obj': page_obj,
    })


@login_required
@role_required('SUPERADMIN', 'ADMIN', 'PANITIA')
def admin_pendaftaran_tambah(request):
    if request.method == 'POST':
        form = AdminPendaftaranForm(request.POST)
        if form.is_valid():
            pendaftaran = form.save()
            messages.success(request, "Pendaftaran berhasil ditambahkan")
            return redirect(
                'admin_pendaftaran_detail',
                nomor=pendaftaran.nomor_pendaftaran
            )
    else:
        form = AdminPendaftaranForm()

    return render(request, 'adminpanel/pendaftaran_form.html', {
        'form': form,
        'mode': 'tambah',
    })


@login_required
@role_required('SUPERADMIN', 'ADMIN')
def admin_pendaftaran_export_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Pendaftaran"

    headers = [
        "No",
        "Nomor Pendaftaran",
        "Nama Lengkap",
        "NIK",
        "Tempat Lahir",
        "Tanggal Lahir",
        "Jenis Kelamin",
        "Agama",
        "Asal Sekolah",
        "Alamat Lengkap",
        "Nama Ayah",
        "Nama Ibu",
        "Pekerjaan Ayah",
        "Pekerjaan Ibu",
        "No WhatsApp",
        "Jurusan",
        "Jalur",
        "Status",
        "NISN",
        "Nilai SKL",
        "Keringanan / Prestasi",
        "Tanggal Daftar",
    ]
    ws.append(headers)

    data = Pendaftaran.objects.all().order_by('tanggal_pendaftaran')

    for idx, p in enumerate(data, start=1):

        alamat_lengkap = (
            f"Dusun {p.dusun or '-'}, "
            f"RT {p.rt or '-'} / RW {p.rw or '-'}, "
            f"Desa {p.desa_kelurahan}, "
            f"Kecamatan {p.kecamatan}, "
            f"{p.kabupaten_kota}"
        )

        ws.append([
            idx,
            p.nomor_pendaftaran,
            p.nama_lengkap,
            p.nik,
            p.tempat_lahir,
            p.tanggal_lahir.strftime("%d-%m-%Y"),
            p.get_jenis_kelamin_display(),
            p.agama,
            p.asal_sekolah,
            alamat_lengkap,
            p.nama_ayah or "-",
            p.nama_ibu or "-",
            p.pekerjaan_ayah or "-",
            p.pekerjaan_ibu or "-",
            p.no_wa,
            p.get_jurusan_display(),
            p.get_jalur_display() if p.jalur else "-",
            p.get_status_display(),
            p.nisn or "-",
            str(p.nilai_skl) if p.nilai_skl else "-",
            p.keringanan_prestasi or "-",
            p.tanggal_pendaftaran.strftime("%d-%m-%Y %H:%M"),
        ])

    # AUTO WIDTH
    for col in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_length + 2

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="data_pendaftaran_lengkap.xlsx"'
    wb.save(response)

    return response

@login_required
@role_required('SUPERADMIN', 'ADMIN')
def admin_pendaftaran_edit(request, pk):
    pendaftaran = get_object_or_404(Pendaftaran, pk=pk)

    if request.method == 'POST':
        form = AdminPendaftaranForm(request.POST, instance=pendaftaran)
        if form.is_valid():
            form.save()
            messages.success(request, "Data pendaftaran berhasil diperbarui.")
            return redirect('admin_pendaftaran_list')
    else:
        form = AdminPendaftaranForm(instance=pendaftaran)

    context = {
        'form': form,
        'mode': 'edit',
        'pendaftaran': pendaftaran,
    }
    return render(request, 'adminpanel/pendaftaran_form.html', context)


# =====================================================
# 🔧 FIX ERROR URL: UBAH STATUS ADMIN
# =====================================================
@login_required
@role_required('SUPERADMIN', 'ADMIN')
@require_POST
def ubah_status_admin(request, pk):
    """
    Mengubah status pendaftaran via admin
    ❗ JIKA SUDAH LUNAS → STATUS TIDAK BOLEH TURUN
    """
    pendaftaran = get_object_or_404(Pendaftaran, pk=pk)
    status_baru = request.POST.get('status')

    if not status_baru:
        messages.error(request, "Status tidak valid.")
        return redirect('admin_pendaftaran_list')

    # 🔒 KUNCI STATUS LUNAS
    if pendaftaran.status == 'lunas':
        messages.warning(
            request,
            "Status sudah LUNAS dan tidak boleh diubah."
        )
        return redirect(
            'admin_pendaftaran_detail',
            nomor=pendaftaran.nomor_pendaftaran
        )

    pendaftaran.status = status_baru
    pendaftaran.save(update_fields=['status'])

    LogAktivitas.objects.create(
        user=request.user,
        pendaftaran=pendaftaran,
        aksi="Ubah Status",
        detail=f"Status diubah menjadi {status_baru}"
    )

    messages.success(request, "Status pendaftaran berhasil diperbarui.")
    return redirect(
        'admin_pendaftaran_detail',
        nomor=pendaftaran.nomor_pendaftaran
    )


# =====================================================
# DETAIL PENDAFTAR
# =====================================================
@login_required
@role_required('SUPERADMIN', 'ADMIN', 'PANITIA', 'BENDAHARA')
def admin_pendaftaran_detail(request, nomor):
    pendaftaran = get_object_or_404(
        Pendaftaran,
        nomor_pendaftaran=nomor
    )

    is_bendahara = request.user.groups.filter(
        name='BENDAHARA'
    ).exists()

    BIAYA_DAFTAR_ULANG = 250000

    total_bayar = pendaftaran.pembayaran_daftar_ulang.aggregate(
        total=Sum('nominal')
    )['total'] or 0

    sisa = BIAYA_DAFTAR_ULANG - total_bayar
    if sisa < 0:
        sisa = 0

    return render(request, 'adminpanel/pendaftaran_detail.html', {
        'pendaftaran': pendaftaran,
        'is_bendahara': is_bendahara,
        'biaya_daftar_ulang': BIAYA_DAFTAR_ULANG,
        'total_bayar': total_bayar,
        'sisa': sisa,
    })


# =====================================================
# DAFTAR ULANG
# =====================================================
@login_required
@role_required('SUPERADMIN', 'ADMIN', 'BENDAHARA')
@require_POST
def bendahara_daftar_ulang(request, pk):
    pendaftaran = get_object_or_404(Pendaftaran, pk=pk)

    # 🔒 BLOK: TERDAFTAR TIDAK BOLEH BAYAR
    if pendaftaran.status == 'terdaftar':
        messages.error(
            request,
            "Pendaftar belum diterima. Tidak bisa melakukan pembayaran."
        )
        return redirect(
            'admin_pendaftaran_detail',
            nomor=pendaftaran.nomor_pendaftaran
        )

    form = PembayaranDaftarUlangForm(request.POST)
    if not form.is_valid():
        messages.error(request, "Data pembayaran tidak valid")
        return redirect(
            'admin_pendaftaran_detail',
            nomor=pendaftaran.nomor_pendaftaran
        )

    pembayaran = form.save(commit=False)
    pembayaran.pendaftaran = pendaftaran
    pembayaran.petugas = request.user
    pembayaran.save()

    BIAYA_DAFTAR_ULANG = 250000

    total_bayar = pendaftaran.pembayaran_daftar_ulang.aggregate(
        total=Sum('nominal')
    )['total'] or 0

    # 🔼 NAIKKAN STATUS (TIDAK PERNAH TURUN)
    if pendaftaran.status == 'diterima':
        pendaftaran.status = 'daftar_ulang'

    if total_bayar >= BIAYA_DAFTAR_ULANG:
        pendaftaran.status = 'lunas'

    pendaftaran.save(update_fields=['status'])

    LogAktivitas.objects.create(
        user=request.user,
        pendaftaran=pendaftaran,
        aksi="Pembayaran Daftar Ulang",
        detail=f"Pembayaran Rp {pembayaran.nominal}"
    )

    messages.success(request, "Pembayaran berhasil disimpan.")
    return redirect(
        'admin_pendaftaran_detail',
        nomor=pendaftaran.nomor_pendaftaran
    )

# =====================================================
# REKAP DAFTAR ULANG + PENCARIAN
# =====================================================
@login_required
@role_required('SUPERADMIN', 'ADMIN', 'BENDAHARA')
def rekap_daftar_ulang(request):
    qs = PembayaranDaftarUlang.objects.select_related(
        'pendaftaran',
        'petugas'
    )

    tanggal_dari = request.GET.get('tanggal_dari')
    tanggal_sampai = request.GET.get('tanggal_sampai')
    petugas = request.GET.get('petugas')
    keyword = request.GET.get('q')

    if tanggal_dari:
        qs = qs.filter(tanggal__gte=tanggal_dari)

    if tanggal_sampai:
        qs = qs.filter(tanggal__lte=tanggal_sampai)

    if petugas:
        qs = qs.filter(petugas_id=petugas)

    if keyword:
        qs = qs.filter(
            Q(pendaftaran__nomor_pendaftaran__icontains=keyword) |
            Q(pendaftaran__nama_lengkap__icontains=keyword) |
            Q(pendaftaran__jurusan__icontains=keyword)
        )

    total_nominal = qs.aggregate(
        total=Sum('nominal')
    )['total'] or 0

    daftar_petugas = (
        qs.values('petugas__id', 'petugas__username')
        .distinct()
    )

    return render(request, 'bendahara/rekap_daftar_ulang.html', {
        'data': qs,
        'total_nominal': total_nominal,
        'daftar_petugas': daftar_petugas,
        'tanggal_dari': tanggal_dari,
        'tanggal_sampai': tanggal_sampai,
        'petugas_selected': petugas,
        'keyword': keyword,
    })


# =====================================================
# EXPORT REKAP
# =====================================================
@login_required
@role_required('SUPERADMIN', 'ADMIN', 'BENDAHARA')
def export_rekap_daftar_ulang_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap Daftar Ulang"

    headers = [
        "No",
        "Nomor Pendaftaran",
        "Nama",
        "Jurusan",
        "Nominal",
        "Tanggal",
        "Petugas",
    ]
    ws.append(headers)

    data = PembayaranDaftarUlang.objects.select_related(
        'pendaftaran',
        'petugas'
    )

    total_nominal = 0

    for idx, p in enumerate(data, start=1):
        ws.append([
            idx,
            p.pendaftaran.nomor_pendaftaran,
            p.pendaftaran.nama_lengkap,
            p.pendaftaran.get_jurusan_display(),
            p.nominal,
            p.tanggal.strftime("%d-%m-%Y"),
            p.petugas.username if p.petugas else "",
        ])
        total_nominal += p.nominal

    ws.append([])
    ws.append(["", "", "", "TOTAL", total_nominal, "", ""])

    for col in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_length + 2

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="rekap_daftar_ulang.xlsx"'
    wb.save(response)

    return response
